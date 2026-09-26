from __future__ import annotations

import os
from pathlib import Path
from dataclasses import dataclass
from difflib import SequenceMatcher
from typing import Optional
from docx import Document
import numpy as np
import faiss
from sentence_transformers import SentenceTransformer
import ebooklib
from ebooklib import epub
from bs4 import BeautifulSoup
from pypdf import PdfReader

# ============================================================
# Embedding model
# ============================================================

# Loaded lazily so importing this file doesn't immediately load
# the embedding model into memory.
_embedding_model = None

EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"


def _get_embedding_model():
    global _embedding_model

    if _embedding_model is None:
        _embedding_model = SentenceTransformer(EMBEDDING_MODEL)

    return _embedding_model


# ============================================================
# Conversation-scoped document memory
# ============================================================

@dataclass
class _IndexedDocument:
    path: str
    mtime_ns: int
    size_bytes: int
    chunk_size: int
    overlap: int
    chunks: list[str]
    index: object


# Temporary RAM-only document memory for the current conversation.
# The FAISS indexes are not added to the LLM context and are not saved
# to disk. Resetting the conversation clears this state.
_document_cache: dict[str, _IndexedDocument] = {}
_active_document: Optional[str] = None


def reset_document_context() -> None:
    """Discard all document indexes associated with the conversation."""

    global _active_document

    _document_cache.clear()
    _active_document = None


# ============================================================
# locate_file
# ============================================================

def locate_file(
    query: str,
    root: Optional[str] = None,
    max_results: int = 10,
) -> str:
    """
    Locate files whose filename or path resembles the requested query.

    Args:
        query:
            Filename, partial filename, extension, or path fragment.
            Examples:
                "shadow of the colossus"
                "research.pdf"
                "farm finance"
                ".epub"

        root:
            Directory to search.

            If omitted, searches the current user's home directory.

            Examples:
                C:\\Users\\Example
                C:\\Documents
                C:\\Data

        max_results:
            Maximum number of matching files returned.

    Returns:
        Human-readable ranked file results.
    """

    if not query.strip():
        return "Error: empty search query."

    query = query.strip()
    query_lower = query.lower()

    if root is None:
        root_path = Path.home()
    else:
        root_path = Path(root).expanduser()

    if not root_path.exists():
        return f"Error: search root does not exist: {root_path}"

    results = []

    def score_path(path: Path) -> float:
        filename = path.name.lower()
        full_path = str(path).lower()

        score = 0.0

        # Exact filename match
        if filename == query_lower:
            score += 100.0

        # Filename without extension exactly matches query
        if path.stem.lower() == query_lower:
            score += 90.0

        # Query appears inside filename
        if query_lower in filename:
            score += 70.0

        # Query appears somewhere in full path
        if query_lower in full_path:
            score += 40.0

        # Fuzzy filename comparison
        similarity = SequenceMatcher(
            None,
            query_lower,
            filename,
        ).ratio()

        score += similarity * 30.0

        return score

    try:
        for current_root, dirs, files in os.walk(
            root_path,
            topdown=True,
            onerror=lambda _: None,
        ):
            # Skip some directories that are generally useless,
            # huge, or inaccessible during ordinary searches.
            dirs[:] = [
                d for d in dirs
                if d.lower() not in {
                    "$recycle.bin",
                    "system volume information",
                    ".git",
                    "__pycache__",
                    "node_modules",
                }
            ]

            for filename in files:
                path = Path(current_root) / filename

                filename_lower = filename.lower()
                full_path_lower = str(path).lower()

                # Cheap prefilter before fuzzy comparison.
                if (
                    query_lower not in filename_lower
                    and query_lower not in full_path_lower
                ):
                    similarity = SequenceMatcher(
                        None,
                        query_lower,
                        filename_lower,
                    ).ratio()

                    # Ignore obviously unrelated filenames.
                    if similarity < 0.35:
                        continue

                score = score_path(path)

                if score > 0:
                    results.append(
                        (
                            score,
                            path,
                        )
                    )

    except Exception as exc:
        return f"File search error: {type(exc).__name__}: {exc}"

    if not results:
        return (
            f"No files matching '{query}' were found under "
            f"{root_path}."
        )

    # Best matches first.
    results.sort(
        key=lambda item: item[0],
        reverse=True,
    )

    results = results[:max_results]

    output = [
        f"Found {len(results)} likely match(es) for '{query}':"
    ]

    for index, (score, path) in enumerate(results, start=1):
        try:
            size = path.stat().st_size
            size_mb = size / (1024 * 1024)
            size_text = f"{size_mb:.2f} MB"
        except OSError:
            size_text = "unknown size"

        output.append(
            f"\n{index}. {path}\n"
            f"   Size: {size_text}\n"
            f"   Match score: {score:.1f}"
        )

    return "\n".join(output)


# ============================================================
# File extraction
# ============================================================

def _extract_pdf(path: Path) -> str:
    

    reader = PdfReader(str(path))

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""

        if text.strip():
            pages.append(
                f"\n--- PAGE {page_number} ---\n{text}"
            )

    return "\n".join(pages)


def _extract_epub(path: Path) -> str:

    book = epub.read_epub(str(path))

    sections = []

    for item in book.get_items():
        if item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue

        soup = BeautifulSoup(
            item.get_content(),
            "html.parser",
        )

        text = soup.get_text(
            separator="\n",
            strip=True,
        )

        if text:
            sections.append(text)

    return "\n\n".join(sections)


def _extract_docx(path: Path) -> str:


    document = Document(str(path))

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n\n".join(paragraphs)


def _extract_plain_text(path: Path) -> str:
    encodings = [
        "utf-8",
        "utf-8-sig",
        "cp1252",
        "latin-1",
    ]

    for encoding in encodings:
        try:
            return path.read_text(
                encoding=encoding,
                errors="strict",
            )
        except UnicodeDecodeError:
            continue

    # Last resort.
    return path.read_text(
        encoding="utf-8",
        errors="replace",
    )


def _extract_text(path: Path) -> str:
    """
    Extract text based on the file extension.
    """

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)

    if suffix == ".epub":
        return _extract_epub(path)

    if suffix == ".docx":
        return _extract_docx(path)

    # Treat these as ordinary text.
    text_extensions = {
        ".txt",
        ".md",
        ".markdown",
        ".py",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".css",
        ".js",
        ".ts",
        ".tsx",
        ".jsx",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
        ".csv",
        ".log",
        ".ps1",
        ".bat",
        ".cmd",
        ".c",
        ".cpp",
        ".h",
        ".hpp",
        ".java",
        ".rs",
        ".go",
        ".sql",
    }

    if suffix in text_extensions:
        return _extract_plain_text(path)

    raise ValueError(
        f"Unsupported file type: {suffix or '[no extension]'}"
    )


# ============================================================
# Chunking
# ============================================================

def _split_text(
    text: str,
    chunk_size: int = 2400,
    overlap: int = 250,
) -> list[str]:
    """
    Split extracted text into overlapping chunks.

    chunk_size and overlap are character counts, not tokens.
    """

    if chunk_size <= 0:
        raise ValueError("chunk_size must be greater than zero.")

    if overlap < 0:
        raise ValueError("overlap cannot be negative.")

    if overlap >= chunk_size:
        raise ValueError(
            "overlap must be smaller than chunk_size."
        )

    # Remove excessive whitespace while preserving paragraph breaks.
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")

    chunks = []

    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(
            start + chunk_size,
            text_length,
        )

        # Try not to cut a paragraph in half.
        if end < text_length:
            paragraph_break = text.rfind(
                "\n\n",
                start,
                end,
            )

            if paragraph_break > start + (chunk_size // 2):
                end = paragraph_break

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= text_length:
            break

        start = max(
            end - overlap,
            start + 1,
        )

    return chunks


# ============================================================
# Conversation-scoped indexing
# ============================================================

def _resolve_file_path(path: str) -> Path:
    return Path(path).expanduser().resolve()


def _cache_is_valid(
    document: _IndexedDocument,
    file_path: Path,
    chunk_size: int,
    overlap: int,
) -> bool:
    stat = file_path.stat()

    return (
        document.mtime_ns == stat.st_mtime_ns
        and document.size_bytes == stat.st_size
        and document.chunk_size == chunk_size
        and document.overlap == overlap
    )


def _load_or_index_document(
    path: str,
    chunk_size: int,
    overlap: int,
) -> tuple[_IndexedDocument, bool]:
    """
    Return (document, newly_indexed).

    The document's chunks and FAISS index remain alive in RAM for the
    current conversation and are reused by later consult_file() calls.
    """

    global _active_document

    file_path = _resolve_file_path(path)

    if not file_path.exists():
        raise FileNotFoundError(f"File does not exist: {file_path}")

    if not file_path.is_file():
        raise ValueError(f"Path is not a file: {file_path}")

    normalized_path = str(file_path)
    cached = _document_cache.get(normalized_path)

    if cached is not None and _cache_is_valid(
        cached,
        file_path,
        chunk_size,
        overlap,
    ):
        _active_document = normalized_path
        return cached, False

    text = _extract_text(file_path)

    if not text.strip():
        raise ValueError("No readable text could be extracted from the file.")

    chunks = _split_text(
        text,
        chunk_size=chunk_size,
        overlap=overlap,
    )

    if not chunks:
        raise ValueError("No searchable chunks were produced.")

    model = _get_embedding_model()

    chunk_embeddings = model.encode(
        chunks,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    chunk_embeddings = np.asarray(
        chunk_embeddings,
        dtype=np.float32,
    )

    # Normalized vectors + inner product = cosine similarity.
    index = faiss.IndexFlatIP(chunk_embeddings.shape[1])
    index.add(chunk_embeddings)

    stat = file_path.stat()

    document = _IndexedDocument(
        path=normalized_path,
        mtime_ns=stat.st_mtime_ns,
        size_bytes=stat.st_size,
        chunk_size=chunk_size,
        overlap=overlap,
        chunks=chunks,
        index=index,
    )

    _document_cache[normalized_path] = document
    _active_document = normalized_path

    return document, True


# ============================================================
# consult_file
# ============================================================

def consult_file(
    path: Optional[str] = None,
    query: str = "",
    top_k: int = 6,
    chunk_size: int = 1800,
    overlap: int = 250,
) -> str:
    """
    Semantically search a local document.

    First call for a document:
        consult_file(
            path=r"C:\\Documents\\book.pdf",
            query="What is this book about?"
        )

    Follow-up calls in the SAME conversation:
        consult_file(
            query="What does it say about travel?"
        )

    On the first call, the file is extracted, chunked, embedded, and placed
    in an in-memory FAISS index. On later calls, that index is reused.

    The index is discarded when reset_document_context() is called, which
    the conversation reset path does when the conversation itself is reset.
    """

    global _active_document

    query = (query or "").strip()

    if not query:
        return "Error: semantic search query is empty."

    # Follow-up questions can omit the path and automatically use the most
    # recently consulted document in this conversation.
    if path is None or not str(path).strip():
        path = _active_document

        if path is None:
            return (
                "Error: no document is currently loaded in this conversation. "
                "Use locate_file if necessary, then call consult_file once "
                "with the document's full path."
            )

    try:
        document, newly_indexed = _load_or_index_document(
            path=str(path),
            chunk_size=chunk_size,
            overlap=overlap,
        )
    except Exception as exc:
        return (
            "Could not consult file:\n"
            f"{type(exc).__name__}: {exc}"
        )

    model = _get_embedding_model()

    try:
        query_embedding = model.encode(
            [query],
            convert_to_numpy=True,
            normalize_embeddings=True,
            show_progress_bar=False,
        )

        query_embedding = np.asarray(
            query_embedding,
            dtype=np.float32,
        )
    except Exception as exc:
        return (
            "Could not create query embedding:\n"
            f"{type(exc).__name__}: {exc}"
        )

    actual_top_k = min(
        max(1, int(top_k)),
        len(document.chunks),
    )

    scores, indices = document.index.search(
        query_embedding,
        actual_top_k,
    )

    status = (
        "Document indexed and loaded for this conversation."
        if newly_indexed
        else "Existing conversation document index reused."
    )

    output = [
        "Semantic search results for:",
        document.path,
        "",
        status,
        f"Query: {query}",
        f"Document chunks indexed: {len(document.chunks)}",
        f"Results returned: {actual_top_k}",
    ]

    for rank, (chunk_index, score) in enumerate(
        zip(indices[0], scores[0]),
        start=1,
    ):
        chunk_index = int(chunk_index)

        if chunk_index < 0:
            continue

        chunk = document.chunks[chunk_index]

        output.append(
            "\n"
            + "=" * 70
            + f"\nRESULT {rank}"
            + f"\nSimilarity: {float(score):.4f}"
            + f"\nChunk: {chunk_index + 1}/{len(document.chunks)}"
            + "\n"
            + "=" * 70
            + "\n"
            + chunk
        )

    return "\n".join(output)
