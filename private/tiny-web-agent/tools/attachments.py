from __future__ import annotations

import json
import re
import uuid
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import numpy as np

from runtime_paths import DATA_DIR
from tools.web_search import _embedding_lock, _get_embedding_model, _cosine_similarity

ATTACHMENTS_DIR = DATA_DIR / "attachments"
MAX_ATTACHMENT_BYTES = 500 * 1024 * 1024
CHUNK_WORDS = 220
CHUNK_OVERLAP = 40
DEFAULT_TOP_K = 4

DOCUMENT_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".html", ".htm",
    ".csv", ".json", ".yaml", ".yml",
}
SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data and data.strip():
            self.parts.append(data.strip())

    def text(self) -> str:
        return "\n".join(self.parts)


def _attachment_dir(attachment_id: str) -> Path:
    clean_id = str(attachment_id or "").strip()
    if not re.fullmatch(r"att_[0-9a-f]{24}", clean_id):
        raise ValueError("Invalid attachment id.")
    return ATTACHMENTS_DIR / clean_id


def _meta_path(attachment_id: str) -> Path:
    return _attachment_dir(attachment_id) / "meta.json"


def load_attachment(attachment_id: str) -> dict[str, Any]:
    path = _meta_path(attachment_id)
    if not path.is_file():
        raise FileNotFoundError(f"Attachment not found: {attachment_id}")
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_attachment_path(attachment_id: str, *, require_kind: str | None = None) -> Path:
    meta = load_attachment(attachment_id)
    if require_kind and meta.get("kind") != require_kind:
        raise ValueError(f"Attachment {attachment_id} is not a {require_kind} attachment.")
    path = _attachment_dir(attachment_id) / str(meta["filename"])
    if not path.is_file():
        raise FileNotFoundError(f"Attachment file is missing: {path}")
    return path.resolve()


def _extract_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF attachments require pypdf. Run: pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(pages), {"pages": len(reader.pages)}


def _extract_docx(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX attachments require python-docx. Run: pip install python-docx") from exc

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts), {"paragraphs": len(document.paragraphs)}


def _extract_document(path: Path) -> tuple[str, dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)

    text = path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(text)
        text = parser.text()
    return text, {}


def _chunk_text(text: str) -> list[str]:
    words = " ".join(str(text or "").split()).split()
    if not words:
        return []
    overlap = min(CHUNK_OVERLAP, CHUNK_WORDS - 1)
    step = CHUNK_WORDS - overlap
    chunks: list[str] = []
    for start in range(0, len(words), step):
        chunk = words[start:start + CHUNK_WORDS]
        if not chunk:
            break
        chunks.append(" ".join(chunk))
        if start + CHUNK_WORDS >= len(words):
            break
    return chunks


def _embed_chunks(chunks: list[str]) -> np.ndarray:
    if not chunks:
        return np.empty((0, 0), dtype=np.float32)
    model = _get_embedding_model()
    with _embedding_lock:
        if hasattr(model, "passage_embed"):
            vectors = list(model.passage_embed(chunks, batch_size=32))
        else:
            vectors = list(model.embed(chunks, batch_size=32))
    return np.asarray(vectors, dtype=np.float32)


def save_attachment(conversation_id: str, upload) -> dict[str, Any]:
    filename = str(upload.filename or "").strip()
    if not filename:
        raise ValueError("Attachment has no filename.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported attachment type '{suffix or '(none)'}'. Supported: {supported}")

    attachment_id = "att_" + uuid.uuid4().hex[:24]
    folder = _attachment_dir(attachment_id)
    folder.mkdir(parents=True, exist_ok=False)
    safe_name = Path(filename).name
    stored = folder / safe_name

    upload.save(stored)
    size = stored.stat().st_size
    if size > MAX_ATTACHMENT_BYTES:
        try:
            stored.unlink(missing_ok=True)
            folder.rmdir()
        except Exception:
            pass
        raise ValueError("Attachment is larger than the 500 MB limit.")

    kind = "document"
    meta: dict[str, Any] = {
        "id": attachment_id,
        "conversation_id": str(conversation_id),
        "filename": safe_name,
        "extension": suffix,
        "kind": kind,
        "size_bytes": size,
        "stored_path": str(stored.resolve()),
        "status": "ready",
    }

    text, extra = _extract_document(stored)
    chunks = _chunk_text(text)
    if not chunks:
        meta["status"] = "no_text"
        meta["note"] = "No extractable text was found in this document."
        meta["chunks"] = 0
    else:
        vectors = _embed_chunks(chunks)
        (folder / "chunks.json").write_text(
            json.dumps(chunks, ensure_ascii=False), encoding="utf-8"
        )
        np.save(folder / "vectors.npy", vectors, allow_pickle=False)
        meta["chunks"] = len(chunks)
    meta.update(extra)

    _meta_path(attachment_id).write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return public_attachment(meta)


def public_attachment(meta: dict[str, Any]) -> dict[str, Any]:
    return {
        key: meta.get(key)
        for key in (
            "id", "filename", "extension", "kind", "size_bytes", "status",
            "chunks", "pages", "paragraphs", "note",
        )
        if meta.get(key) is not None
    }


def validate_attachment_ids(conversation_id: str, attachment_ids: list[str]) -> list[dict[str, Any]]:
    metas: list[dict[str, Any]] = []
    seen: set[str] = set()
    for attachment_id in attachment_ids:
        attachment_id = str(attachment_id or "").strip()
        if not attachment_id or attachment_id in seen:
            continue
        seen.add(attachment_id)
        meta = load_attachment(attachment_id)
        if str(meta.get("conversation_id")) != str(conversation_id):
            raise PermissionError(f"Attachment does not belong to this conversation: {attachment_id}")
        metas.append(meta)
    return metas


def delete_conversation_attachments(conversation_ids: list[str]) -> int:
    """Delete stored attachment folders owned by deleted conversations."""
    import shutil

    wanted = {str(item) for item in conversation_ids}
    if not wanted or not ATTACHMENTS_DIR.is_dir():
        return 0
    deleted = 0
    for folder in ATTACHMENTS_DIR.iterdir():
        if not folder.is_dir():
            continue
        meta_path = folder / "meta.json"
        try:
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if str(meta.get("conversation_id")) in wanted:
            shutil.rmtree(folder, ignore_errors=True)
            deleted += 1
    return deleted


def search_attachment(attachment_id: str, query: str, top_k: int = DEFAULT_TOP_K) -> str:
    meta = load_attachment(attachment_id)
    if meta.get("kind") != "document":
        raise ValueError("search_attachment only works with document attachments.")
    if meta.get("status") != "ready":
        return str(meta.get("note") or "This attachment has no searchable text.")

    folder = _attachment_dir(attachment_id)
    chunks = json.loads((folder / "chunks.json").read_text(encoding="utf-8"))
    vectors = np.load(folder / "vectors.npy", allow_pickle=False)
    if not chunks or vectors.size == 0:
        return "No searchable text chunks are available."

    model = _get_embedding_model()
    with _embedding_lock:
        if hasattr(model, "query_embed"):
            query_vector = next(model.query_embed([str(query)]))
        else:
            query_vector = next(model.embed([str(query)], batch_size=1))

    ranked = [
        (_cosine_similarity(query_vector, vector), index)
        for index, vector in enumerate(vectors)
    ]
    ranked.sort(key=lambda item: item[0], reverse=True)
    limit = max(1, min(int(top_k or DEFAULT_TOP_K), 10, len(ranked)))

    parts = [f"Attachment: {meta['filename']} ({attachment_id})"]
    for rank, (score, index) in enumerate(ranked[:limit], start=1):
        parts.append(
            f"Relevant excerpt {rank} (semantic similarity {score:.3f}):\n{chunks[index]}"
        )
    return "\n\n".join(parts)


def build_attachment_context(conversation_id: str, attachment_ids: list[str], query: str) -> str:
    metas = validate_attachment_ids(conversation_id, attachment_ids)
    if not metas:
        return ""

    parts = [
        "ATTACHED FILES (managed by Tiny Web Agent):",
        "Use attachment IDs in tools; do not invent local paths.",
    ]
    for meta in metas:
        attachment_id = meta["id"]
        parts.append(
            f'- {attachment_id}: "{meta["filename"]}" [document, {meta.get("chunks", 0)} chunks].'
        )
        if meta.get("status") == "ready":
            try:
                parts.append(search_attachment(attachment_id, query, top_k=DEFAULT_TOP_K))
            except Exception as error:
                parts.append(f"Automatic retrieval failed for {attachment_id}: {type(error).__name__}: {error}")
        else:
            parts.append(str(meta.get("note") or "No extractable text was found."))

    return "\n\n".join(parts)
